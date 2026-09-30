"""Generate small, valid sample files used as uploads by the monitor."""
import os
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
from PIL import Image, ImageDraw
from docx import Document
from openpyxl import Workbook
from pptx import Presentation
from pptx.util import Inches

D = os.path.join(os.path.dirname(__file__), "..", "fixtures")
os.makedirs(D, exist_ok=True)
p = lambda n: os.path.join(D, n)

def pdf(name, pages=3, label="FileMagics monitor"):
    c = canvas.Canvas(p(name), pagesize=letter)
    for i in range(1, pages + 1):
        c.setFont("Helvetica-Bold", 28); c.drawString(72, 700, f"{label} - page {i}")
        c.setFont("Helvetica", 14)
        c.drawString(72, 660, "The quick brown fox jumps over the lazy dog. 0123456789")
        c.rect(72, 400, 300, 150); c.drawString(90, 470, "Table cell A | Table cell B")
        c.showPage()
    c.save()

pdf("sample.pdf", 3)
pdf("sample2.pdf", 2, "Second file")

img = Image.new("RGB", (400, 300), "white"); d = ImageDraw.Draw(img)
d.rectangle([40, 40, 360, 260], outline="black", width=4); d.ellipse([140, 90, 260, 210], fill="#2e7d32")
d.text((60, 60), "FileMagics OCR test 123", fill="black")
img.save(p("sample.png")); img.save(p("sample.jpg"), quality=90); img.save(p("sample.webp"))
img.save(p("sample2.jpg"), quality=90)

open(p("sample.svg"), "w").write('<svg xmlns="http://www.w3.org/2000/svg" width="200" height="200"><rect width="200" height="200" fill="#fff"/><circle cx="100" cy="100" r="60" fill="#2e7d32"/></svg>')
open(p("sample.html"), "w").write("<!doctype html><html><head><meta charset='utf-8'><title>Test</title></head><body><h1>FileMagics monitor</h1><p>Hello from the daily test.</p></body></html>")
open(p("sample.md"), "w").write("# FileMagics monitor\n\nThis is a **daily** test.\n\n- item one\n- item two\n")

doc = Document(); doc.add_heading("FileMagics monitor", 0); doc.add_paragraph("Daily Word to PDF test."); doc.save(p("sample.docx"))
wb = Workbook(); ws = wb.active; ws.append(["Tool", "Status"]); ws.append(["excel-to-pdf", "ok"]); wb.save(p("sample.xlsx"))
pr = Presentation(); s = pr.slides.add_slide(pr.slide_layouts[1]); s.shapes.title.text = "FileMagics monitor"; s.placeholders[1].text = "Daily PPT to PDF test"; pr.save(p("sample.pptx"))
print("fixtures ok:", sorted(os.listdir(D)))
